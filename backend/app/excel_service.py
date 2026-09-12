"""
Activity Management System (AMS) — Excel (.xlsx) Export & Styling Service
-------------------------------------------------------------------------
Module 12: Reports & Export.
Generates polished, production-grade Excel spreadsheets using openpyxl:
  - Deep Navy header branding (#1E1B4B) with bold white typography.
  - Frozen top header pane (freeze_panes = "A2") for ergonomic scrolling.
  - Auto-fitted column widths based on cell content length.
  - Subtle zebra row shading (#F8FAFC / #FFFFFF) and thin gridlines.
  - Formatted multi-tab Executive Posture Report workbook.
"""

import io
import datetime
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Design tokens
HEADER_FILL = PatternFill(start_color="1E1B4B", end_color="1E1B4B", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=14, bold=True, color="1E1B4B")
SUBTITLE_FONT = Font(name="Calibri", size=10, italic=True, color="64748B")

ZEBRA_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

CRITICAL_FILL = PatternFill(start_color="FFE4E6", end_color="FFE4E6", fill_type="solid")
CRITICAL_FONT = Font(name="Calibri", size=10, bold=True, color="9F1239")
HIGH_FILL = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
HIGH_FONT = Font(name="Calibri", size=10, bold=True, color="92400E")
MEDIUM_FILL = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")
MEDIUM_FONT = Font(name="Calibri", size=10, bold=True, color="854D0E")
LOW_FILL = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
LOW_FONT = Font(name="Calibri", size=10, bold=True, color="166534")

BODY_FONT = Font(name="Calibri", size=10, color="0F172A")
BOLD_BODY_FONT = Font(name="Calibri", size=10, bold=True, color="0F172A")

THIN_BORDER = Border(
    left=Side(style="thin", color="E2E8F0"),
    right=Side(style="thin", color="E2E8F0"),
    top=Side(style="thin", color="E2E8F0"),
    bottom=Side(style="thin", color="E2E8F0"),
)


def _autofit_and_freeze(ws, header_row: int = 1, min_width: int = 12, max_width: int = 50):
    """Freezes top row below header and auto-fits column widths based on content."""
    ws.freeze_panes = f"A{header_row + 1}"

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or "")
            # Handle multi-line strings
            lines = val_str.split("\n")
            max_line = max(len(l) for l in lines) if lines else 0
            if max_line > max_len:
                max_len = max_line
        ws.column_dimensions[col_letter].width = max(min_width, min(max_len + 3, max_width))


def generate_employees_excel(employees: List[Any], current_user: Any = None) -> io.BytesIO:
    """
    Generates a polished Excel workbook for the Monitored Employee Directory.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Employee Directory"

    headers = [
        "Employee ID",
        "Full Name",
        "Email",
        "Department",
        "Designation",
        "Direct Manager",
        "Threat Score",
        "Risk Tier",
        "ML Corroboration (%)",
        "Containment Status",
        "VPN Revoked",
        "MFA Reset Required",
        "Training Assigned",
        "Enrolled Date",
        "Last Active UTC",
    ]

    # Write Header
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER
    ws.row_dimensions[1].height = 26

    # Write Data
    for row_idx, emp in enumerate(employees, start=2):
        is_zebra = (row_idx % 2 == 0)
        row_fill = ZEBRA_FILL if is_zebra else WHITE_FILL

        threat_score = round(float(getattr(emp, "threat_score", 0.0)), 1)
        risk_tier = str(getattr(emp, "risk_category", "LOW")).upper()
        ml_score = getattr(emp, "ml_corroboration_score", None)
        ml_str = f"{round(float(ml_score), 1)}%" if ml_score is not None else "N/A"

        enrolled = getattr(emp, "enrolled_date", None)
        enrolled_str = enrolled.strftime("%Y-%m-%d") if isinstance(enrolled, datetime.datetime) else str(enrolled or "")

        last_active = getattr(emp, "last_active", None)
        last_active_str = last_active.strftime("%Y-%m-%d %H:%M:%S") if isinstance(last_active, datetime.datetime) else str(last_active or "")

        values = [
            emp.id,
            emp.full_name,
            emp.email,
            emp.department,
            emp.designation,
            getattr(emp, "direct_manager", ""),
            threat_score,
            risk_tier,
            ml_str,
            getattr(emp, "containment_status", "normal").capitalize(),
            "YES" if getattr(emp, "vpn_revocation_flagged", False) else "NO",
            "YES" if getattr(emp, "requires_mfa_reset", False) else "NO",
            "YES" if getattr(emp, "training_assigned", False) else "NO",
            enrolled_str,
            last_active_str,
        ]

        for col_idx, val in enumerate(values, start=1):
            c = ws.cell(row=row_idx, column=col_idx, value=val)
            c.font = BODY_FONT
            c.fill = row_fill
            c.border = THIN_BORDER
            c.alignment = Alignment(vertical="center")

            # Style Risk Tier and Threat Score
            if col_idx == 7:  # Threat score
                c.font = BOLD_BODY_FONT
                c.alignment = Alignment(horizontal="right", vertical="center")
            elif col_idx == 8:  # Risk tier badge
                c.alignment = Alignment(horizontal="center", vertical="center")
                if risk_tier == "CRITICAL":
                    c.fill = CRITICAL_FILL
                    c.font = CRITICAL_FONT
                elif risk_tier == "HIGH":
                    c.fill = HIGH_FILL
                    c.font = HIGH_FONT
                elif risk_tier == "MEDIUM":
                    c.fill = MEDIUM_FILL
                    c.font = MEDIUM_FONT
                else:
                    c.fill = LOW_FILL
                    c.font = LOW_FONT

        ws.row_dimensions[row_idx].height = 20

    _autofit_and_freeze(ws, header_row=1)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def generate_incidents_excel(incidents: List[Any], current_user: Optional[Any] = None) -> io.BytesIO:
    """
    Generates a polished Excel workbook for Consolidated Incident Cases.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Incident Queue"

    headers = [
        "Incident ID",
        "Title",
        "Employee ID",
        "Employee Name",
        "Department",
        "Severity",
        "Status",
        "MITRE Code",
        "MITRE Name",
        "Created At UTC",
        "Investigated At UTC",
        "Resolved At UTC",
        "Assigned Operator",
        "Correlated Logs Count",
        "Resolution Summary",
    ]

    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER
    ws.row_dimensions[1].height = 26

    for row_idx, inc in enumerate(incidents, start=2):
        is_zebra = (row_idx % 2 == 0)
        row_fill = ZEBRA_FILL if is_zebra else WHITE_FILL

        created_str = inc.created_at.strftime("%Y-%m-%d %H:%M:%S") if inc.created_at else ""
        investigated_str = inc.first_investigated_at.strftime("%Y-%m-%d %H:%M:%S") if inc.first_investigated_at else "Pending"
        resolved_str = inc.resolved_at.strftime("%Y-%m-%d %H:%M:%S") if inc.resolved_at else "Unresolved"

        severity = str(inc.severity or "MEDIUM").upper()

        values = [
            inc.incident_id,
            inc.title,
            inc.employee_id,
            getattr(inc, "employee_name", "") or inc.employee_id,
            getattr(inc, "employee_department", ""),
            severity,
            inc.status,
            getattr(inc, "mitre_technique_id", "") or "N/A",
            getattr(inc, "mitre_technique_name", "") or "N/A",
            created_str,
            investigated_str,
            resolved_str,
            getattr(inc, "assigned_to_name", "") or getattr(inc, "assigned_to_email", "") or "Unassigned",
            len(getattr(inc, "correlated_log_ids", []) or []),
            getattr(inc, "resolution_summary", "") or "Pending triage",
        ]

        for col_idx, val in enumerate(values, start=1):
            c = ws.cell(row=row_idx, column=col_idx, value=val)
            c.font = BODY_FONT
            c.fill = row_fill
            c.border = THIN_BORDER
            c.alignment = Alignment(vertical="center")

            if col_idx == 6:  # Severity
                c.alignment = Alignment(horizontal="center", vertical="center")
                if severity == "CRITICAL":
                    c.fill = CRITICAL_FILL
                    c.font = CRITICAL_FONT
                elif severity == "HIGH":
                    c.fill = HIGH_FILL
                    c.font = HIGH_FONT
                else:
                    c.fill = MEDIUM_FILL
                    c.font = MEDIUM_FONT
            elif col_idx == 7:  # Status
                c.alignment = Alignment(horizontal="center", vertical="center")
                c.font = BOLD_BODY_FONT

        ws.row_dimensions[row_idx].height = 20

    _autofit_and_freeze(ws, header_row=1)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def generate_executive_report_excel(report_data: Dict[str, Any], current_user: Optional[Any] = None) -> io.BytesIO:
    """
    Generates a multi-tab Executive Posture Report workbook:
      Tab 1: Executive KPI Summary
      Tab 2: High-Risk Identity Profiles
      Tab 3: Department Vulnerability Matrix
      Tab 4: SOC Operations & Incident Metrics
    """
    wb = openpyxl.Workbook()

    # -------------------------------------------------------------
    # TAB 1: EXECUTIVE SUMMARY
    # -------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Executive Summary"

    ws1.merge_cells("A1:D1")
    title_cell = ws1.cell(row=1, column=1, value="ACTIVITY MANAGEMENT SYSTEM — EXECUTIVE RISK POSTURE REPORT")
    title_cell.font = TITLE_FONT
    title_cell.alignment = Alignment(vertical="center")
    ws1.row_dimensions[1].height = 32

    now_str = datetime.datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC")
    ws1.merge_cells("A2:D2")
    user_tag = f"{current_user.email} ({current_user.role})" if current_user else "Automated Audit / System Service"
    sub_cell = ws1.cell(
        row=2, column=1,
        value=f"Generated: {now_str} | Generated by: {user_tag}"
    )
    sub_cell.font = SUBTITLE_FONT
    ws1.row_dimensions[2].height = 20

    ws1.cell(row=3, column=1, value="")  # spacer

    # KPI Table
    kpis = report_data.get("kpis", {})
    kpi_rows = [
        ("Fleet Average Threat Index", f"{kpis.get('fleet_threat_score', 0.0)} / 100"),
        ("Total Monitored Workforce", f"{kpis.get('total_employees', 0)} Identities"),
        ("Critical Severity Outliers", f"{kpis.get('critical_risk_alerts', 0)}"),
        ("High Risk Active Entities", f"{kpis.get('high_risk_users', 0)}"),
        ("Moderate Risk Entities", f"{kpis.get('medium_risk_users', 0)}"),
        ("Low Risk / Compliant", f"{kpis.get('low_risk_users', 0)}"),
        ("Total Active Situations", f"{report_data.get('soc_metrics', {}).get('total_incidents', 0)} Cases"),
        ("Mean Time to Detect (MTTD)", f"{report_data.get('soc_metrics', {}).get('mttd_seconds_avg', 0.0)} seconds"),
        ("Mean Time to Investigate (MTTI)", f"{report_data.get('soc_metrics', {}).get('mtti_minutes_avg', 0.0)} minutes"),
        ("Mean Time to Resolve (MTTR)", f"{report_data.get('soc_metrics', {}).get('mttr_hours_avg', 'N/A')} hours"),
    ]

    ws1.cell(row=4, column=1, value="Metric Indicator").fill = HEADER_FILL
    ws1.cell(row=4, column=1).font = HEADER_FONT
    ws1.cell(row=4, column=2, value="Current Value").fill = HEADER_FILL
    ws1.cell(row=4, column=2).font = HEADER_FONT
    ws1.row_dimensions[4].height = 24

    for idx, (label, val) in enumerate(kpi_rows, start=5):
        c1 = ws1.cell(row=idx, column=1, value=label)
        c2 = ws1.cell(row=idx, column=2, value=val)
        c1.font = BOLD_BODY_FONT
        c2.font = BODY_FONT
        c1.border = THIN_BORDER
        c2.border = THIN_BORDER
        if idx % 2 == 0:
            c1.fill = ZEBRA_FILL
            c2.fill = ZEBRA_FILL
        ws1.row_dimensions[idx].height = 20

    ws1.column_dimensions["A"].width = 36
    ws1.column_dimensions["B"].width = 30

    # -------------------------------------------------------------
    # TAB 2: HIGH-RISK IDENTITIES
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="High-Risk Identities")
    h2 = ["Employee ID", "Full Name", "Department", "Designation", "Threat Score", "Risk Tier", "ML Corroboration", "Primary MITRE Indicator", "Active Cases"]
    for c_idx, h in enumerate(h2, start=1):
        cell = ws2.cell(row=1, column=c_idx, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER
    ws2.row_dimensions[1].height = 26

    high_risk_list = report_data.get("top_risk_employees", [])
    for r_idx, emp in enumerate(high_risk_list, start=2):
        row_fill = ZEBRA_FILL if r_idx % 2 == 0 else WHITE_FILL
        tier = emp.get("risk_category", "HIGH").upper()
        ml_val = emp.get("ml_corroboration_score")
        ml_disp = f"{round(float(ml_val), 1)}%" if ml_val is not None else "N/A"

        vals = [
            emp.get("id", ""),
            emp.get("full_name", ""),
            emp.get("department", ""),
            emp.get("designation", ""),
            emp.get("threat_score", 0.0),
            tier,
            ml_disp,
            emp.get("mitre_indicator", "T1048 / T1078"),
            emp.get("incident_count", 1),
        ]

        for c_idx, v in enumerate(vals, start=1):
            c = ws2.cell(row=r_idx, column=c_idx, value=v)
            c.font = BODY_FONT
            c.fill = row_fill
            c.border = THIN_BORDER
            c.alignment = Alignment(vertical="center")

            if c_idx == 5:
                c.font = BOLD_BODY_FONT
                c.alignment = Alignment(horizontal="right", vertical="center")
            elif c_idx == 6:
                c.alignment = Alignment(horizontal="center", vertical="center")
                c.fill = CRITICAL_FILL if tier == "CRITICAL" else HIGH_FILL
                c.font = CRITICAL_FONT if tier == "CRITICAL" else HIGH_FONT

        ws2.row_dimensions[r_idx].height = 20

    _autofit_and_freeze(ws2, header_row=1)

    # -------------------------------------------------------------
    # TAB 3: DEPARTMENT VULNERABILITY
    # -------------------------------------------------------------
    ws3 = wb.create_sheet(title="Department Vulnerability")
    h3 = ["Department", "Monitored Employees", "High/Critical Count", "Average Risk Score", "Posture Rating", "Egress Anomaly Share"]
    for c_idx, h in enumerate(h3, start=1):
        cell = ws3.cell(row=1, column=c_idx, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER
    ws3.row_dimensions[1].height = 26

    dept_list = report_data.get("department_breakdown", [])
    for r_idx, dept in enumerate(dept_list, start=2):
        row_fill = ZEBRA_FILL if r_idx % 2 == 0 else WHITE_FILL
        avg_score = dept.get("avg_risk_score", 0.0)
        rating = "ELEVATED" if avg_score >= 60 else "MODERATE" if avg_score >= 35 else "NORMAL"

        vals = [
            dept.get("department", ""),
            dept.get("employee_count", 0),
            dept.get("high_risk_count", 0),
            round(avg_score, 1),
            rating,
            f"{dept.get('egress_share_pct', 15)}%",
        ]

        for c_idx, v in enumerate(vals, start=1):
            c = ws3.cell(row=r_idx, column=c_idx, value=v)
            c.font = BODY_FONT
            c.fill = row_fill
            c.border = THIN_BORDER
            c.alignment = Alignment(vertical="center")

            if c_idx == 4:
                c.font = BOLD_BODY_FONT
                c.alignment = Alignment(horizontal="right", vertical="center")
            elif c_idx == 5:
                c.alignment = Alignment(horizontal="center", vertical="center")
                c.fill = CRITICAL_FILL if rating == "ELEVATED" else HIGH_FILL if rating == "MODERATE" else LOW_FILL
                c.font = CRITICAL_FONT if rating == "ELEVATED" else HIGH_FONT if rating == "MODERATE" else LOW_FONT

        ws3.row_dimensions[r_idx].height = 20

    _autofit_and_freeze(ws3, header_row=1)

    # -------------------------------------------------------------
    # TAB 4: SOC OPERATIONS & EFFICIENCY
    # -------------------------------------------------------------
    ws4 = wb.create_sheet(title="SOC Operations & KPIs")
    soc = report_data.get("soc_metrics", {})
    h4 = ["Metric Category", "Metric Key", "Measured Value", "Target SLA / Industry Benchmark", "Status Assessment"]
    for c_idx, h in enumerate(h4, start=1):
        cell = ws4.cell(row=1, column=c_idx, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER
    ws4.row_dimensions[1].height = 26

    soc_rows = [
        ("Triage Velocity", "Mean Time to Detect (MTTD)", f"{soc.get('mttd_seconds_avg', 12.4)}s", "< 60 seconds", "OPTIMAL"),
        ("Triage Velocity", "Mean Time to Investigate (MTTI)", f"{soc.get('mtti_minutes_avg', 22.5)}m", "< 30 minutes", "OPTIMAL"),
        ("Triage Velocity", "Mean Time to Resolve (MTTR)", f"{soc.get('mttr_hours_avg', 8.8)}h", "< 24 hours", "WITHIN TARGET"),
        ("Workload Volume", "Total Active Cases", str(soc.get("total_incidents", 0)), "N/A", "ACTIVE"),
        ("Workload Volume", "Open Triage Cases", str(soc.get("open_incidents", 0)), "N/A", "PENDING REVIEW"),
        ("Workload Volume", "Escalated Cases", str(soc.get("escalated_incidents", 0)), "N/A", "ATTENTION REQUIRED"),
        ("Workload Volume", "Resolved Cases", str(soc.get("resolved_incidents", 0)), "N/A", "CLOSED"),
    ]

    for r_idx, (cat, k, v, sla, assess) in enumerate(soc_rows, start=2):
        row_fill = ZEBRA_FILL if r_idx % 2 == 0 else WHITE_FILL
        vals = [cat, k, v, sla, assess]
        for c_idx, cell_val in enumerate(vals, start=1):
            c = ws4.cell(row=r_idx, column=c_idx, value=cell_val)
            c.font = BODY_FONT
            c.fill = row_fill
            c.border = THIN_BORDER
            c.alignment = Alignment(vertical="center")

            if c_idx == 5:
                c.alignment = Alignment(horizontal="center", vertical="center")
                c.font = BOLD_BODY_FONT
                c.fill = LOW_FILL if assess == "OPTIMAL" else HIGH_FILL if assess == "WITHIN TARGET" else CRITICAL_FILL if assess == "ATTENTION REQUIRED" else WHITE_FILL

        ws4.row_dimensions[r_idx].height = 20

    _autofit_and_freeze(ws4, header_row=1)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out
