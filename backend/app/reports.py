from io import BytesIO
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib import colors
from openpyxl import Workbook
from sqlalchemy.orm import Session
from .models import Employee, RiskScore, Anomaly, Incident

def make_pdf(db: Session):
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4)
    story = [Paragraph("Insider Threat Behavioral Intelligence Report", style=None), Spacer(1, 12)]
    rows = [["Employee", "Department", "Risk Score", "Category"]]
    for e in db.query(Employee).all():
        r = db.query(RiskScore).filter(RiskScore.employee_id == e.employee_id).first()
        rows.append([e.name, e.department, f"{r.total_score:.2f}" if r else "0", r.category if r else "Low Risk"])
    table = Table(rows)
    table.setStyle(TableStyle([
        ("GRID",(0,0),(-1,-1),0.5,colors.grey),
        ("BACKGROUND",(0,0),(-1,0),colors.lightgrey),
        ("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),
        ("VALIGN",(0,0),(-1,-1),"TOP")
    ]))
    story.append(table)
    story.append(Spacer(1, 18))
    story.append(Paragraph(f"Open anomalies: {db.query(Anomaly).filter(Anomaly.status!='Closed').count()}", style=None))
    story.append(Paragraph(f"Incidents: {db.query(Incident).count()}", style=None))
    doc.build(story)
    buf.seek(0)
    return buf

def make_excel(db: Session):
    wb = Workbook()
    ws = wb.active; ws.title = "Risk Scores"
    ws.append(["Employee", "Department", "Score", "Category"])
    for e in db.query(Employee).all():
        r = db.query(RiskScore).filter(RiskScore.employee_id == e.employee_id).first()
        ws.append([e.name, e.department, r.total_score if r else 0, r.category if r else "Low Risk"])
    wa = wb.create_sheet("Anomalies")
    wa.append(["ID","Employee","Category","Severity","Score","Status","Created"])
    for a in db.query(Anomaly).all():
        wa.append([a.id,a.employee_id,a.category,a.severity,a.score,a.status,a.created_at.isoformat()])
    wi = wb.create_sheet("Incidents")
    wi.append(["ID","Employee","Title","Severity","Status","Assignee","Created"])
    for i in db.query(Incident).all():
        wi.append([i.id,i.employee_id,i.title,i.severity,i.status,i.assignee,i.created_at.isoformat()])
    buf = BytesIO(); wb.save(buf); buf.seek(0)
    return buf
