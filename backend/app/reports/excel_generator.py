"""Excel report generators using openpyxl.

Each public function accepts a SQLAlchemy *Session*, queries the database,
and writes a formatted .xlsx workbook to *output_path*.

House style:
  - Deep-navy header row with white bold text
  - Alternating row shading (white / light blue-grey)
  - Auto-fit column widths (with a max cap to keep columns readable)
  - Severity / risk-band colour coding
  - Freeze panes on row 1 (header)
  - Metadata sheet with generation timestamp
"""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    GradientFill,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models.activity import ActivityLog, ActivityType
from app.models.alert import Alert, AlertStatus
from app.models.anomaly import Anomaly, AnomalyCategory, AnomalyStatus, Severity
from app.models.behavior import BehaviorProfile
from app.models.employee import Employee
from app.models.investigation import Investigation, InvestigationStatus
from app.models.risk import RiskBand, RiskScore

# ---------------------------------------------------------------------------
# Palette (hex without leading #)
# ---------------------------------------------------------------------------
_NAVY = "0F172A"
_INDIGO = "4F46E5"
_TEAL = "0D9488"
_LIGHT_GREY = "F1F5F9"
_WHITE = "FFFFFF"
_ALT_ROW = "EFF6FF"

_SEVERITY_HEX: dict[str, str] = {
    Severity.CRITICAL: "DC2626",
    Severity.HIGH: "EA580C",
    Severity.MEDIUM: "D97706",
    Severity.LOW: "65A30D",
    Severity.INFORMATIONAL: "0EA5E9",
}

_RISK_HEX: dict[str, str] = {
    RiskBand.CRITICAL: "DC2626",
    RiskBand.HIGH: "EA580C",
    RiskBand.MEDIUM: "D97706",
    RiskBand.LOW: "65A30D",
}

_ALERT_STATUS_HEX: dict[str, str] = {
    AlertStatus.OPEN: "DC2626",
    AlertStatus.ACKNOWLEDGED: "D97706",
    AlertStatus.CLOSED: "65A30D",
}

_INV_STATUS_HEX: dict[str, str] = {
    InvestigationStatus.OPEN: "DC2626",
    InvestigationStatus.IN_PROGRESS: "D97706",
    InvestigationStatus.ESCALATED: "EA580C",
    InvestigationStatus.RESOLVED: "65A30D",
    InvestigationStatus.CLOSED: "94A3B8",
}

# ---------------------------------------------------------------------------
# Shared style helpers
# ---------------------------------------------------------------------------


def _fill(hex_color: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=hex_color)


def _thin_border() -> Border:
    side = Side(style="thin", color="CBD5E1")
    return Border(left=side, right=side, top=side, bottom=side)


def _header_font() -> Font:
    return Font(name="Calibri", bold=True, color=_WHITE, size=10)


def _body_font(bold: bool = False) -> Font:
    return Font(name="Calibri", bold=bold, size=9)


def _center_align() -> Alignment:
    return Alignment(horizontal="center", vertical="center", wrap_text=True)


def _left_align() -> Alignment:
    return Alignment(horizontal="left", vertical="center", wrap_text=True)


def _apply_header_row(ws, row: int, headers: list[str]) -> None:
    """Write a navy header row with white bold text."""
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col_idx, value=header)
        cell.font = _header_font()
        cell.fill = _fill(_NAVY)
        cell.alignment = _center_align()
        cell.border = _thin_border()
    ws.row_dimensions[row].height = 20


def _apply_data_row(
    ws,
    row: int,
    values: list,
    alternate: bool,
    color_map: dict[int, str] | None = None,
) -> None:
    """Write a data row with optional alternating shading and per-cell colour overrides.

    Parameters
    ----------
    color_map : {col_index (1-based): hex_color_string} — overrides background for that cell
    """
    base_fill = _fill(_ALT_ROW if alternate else _WHITE)
    border = _thin_border()
    for col_idx, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col_idx, value=value)
        cell.font = _body_font()
        cell.fill = base_fill
        cell.border = border
        if col_idx <= 2:
            cell.alignment = _left_align()
        else:
            cell.alignment = _center_align()

    # Apply colour overrides (badge-style cells)
    if color_map:
        for col_idx, hex_color in color_map.items():
            cell = ws.cell(row=row, column=col_idx)
            cell.fill = _fill(hex_color)
            cell.font = Font(name="Calibri", bold=True, color=_WHITE, size=9)
            cell.alignment = _center_align()


def _auto_column_widths(ws, min_width: int = 8, max_width: int = 45) -> None:
    """Iterate over all columns and set width based on max cell content length."""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, min_width), max_width)


def _freeze_header(ws) -> None:
    ws.freeze_panes = "A2"


def _add_metadata_sheet(wb: Workbook, report_name: str) -> None:
    ws = wb.create_sheet("Metadata")
    rows = [
        ("Report Name", report_name),
        ("Generated At", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")),
        ("Platform", "InsiderIQ"),
        ("Classification", "Confidential"),
    ]
    _apply_header_row(ws, 1, ["Property", "Value"])
    for r, (prop, val) in enumerate(rows, start=2):
        ws.cell(row=r, column=1, value=prop).font = _body_font(bold=True)
        ws.cell(row=r, column=2, value=val).font = _body_font()
        ws.cell(row=r, column=1).border = _thin_border()
        ws.cell(row=r, column=2).border = _thin_border()
    _auto_column_widths(ws)


def _now_str() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")


def _latest_rs_subq():
    """Subquery: latest computed_at per employee for risk scores."""
    return (
        select(
            RiskScore.employee_id,
            func.max(RiskScore.computed_at).label("max_computed_at"),
        )
        .group_by(RiskScore.employee_id)
        .subquery()
    )


# ---------------------------------------------------------------------------
# 1. Insider Threat Excel Report
# ---------------------------------------------------------------------------


def generate_insider_threat_excel(db: Session, output_path: str | Path) -> Path:
    """Generate a multi-sheet insider threat Excel workbook.

    Sheets
    ------
    - Summary        : KPI statistics
    - Anomalies      : All anomaly records
    - Alerts         : All alert records
    - Investigations : All investigation records
    - Risk Scores    : Latest risk score per employee
    - Metadata
    """
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    wb.remove(wb.active)  # remove default sheet

    # ── Sheet 1: Summary ────────────────────────────────────────────────────
    ws_sum = wb.create_sheet("Summary")
    _apply_header_row(ws_sum, 1, ["Metric", "Value"])

    total_employees = db.scalar(select(func.count()).select_from(Employee)) or 0
    total_anomalies = db.scalar(select(func.count()).select_from(Anomaly)) or 0
    open_alerts = (
        db.scalar(select(func.count()).select_from(Alert).where(Alert.status == AlertStatus.OPEN))
        or 0
    )
    open_invs = (
        db.scalar(
            select(func.count())
            .select_from(Investigation)
            .where(
                Investigation.status.in_(
                    [InvestigationStatus.OPEN, InvestigationStatus.IN_PROGRESS]
                )
            )
        )
        or 0
    )
    critical_anomalies = (
        db.scalar(
            select(func.count()).select_from(Anomaly).where(Anomaly.severity == Severity.CRITICAL)
        )
        or 0
    )
    total_alerts = db.scalar(select(func.count()).select_from(Alert)) or 0
    total_invs = db.scalar(select(func.count()).select_from(Investigation)) or 0

    summary_data = [
        ("Total Monitored Employees", total_employees),
        ("Total Anomalies Detected", total_anomalies),
        ("Critical Anomalies", critical_anomalies),
        ("Total Alerts", total_alerts),
        ("Open Alerts", open_alerts),
        ("Total Investigations", total_invs),
        ("Active Investigations", open_invs),
        ("Report Generated", _now_str()),
    ]
    for r, (metric, val) in enumerate(summary_data, start=2):
        ws_sum.cell(row=r, column=1, value=metric).font = _body_font(bold=True)
        ws_sum.cell(row=r, column=2, value=val).font = _body_font()
        ws_sum.cell(row=r, column=1).fill = _fill(_LIGHT_GREY if r % 2 == 0 else _WHITE)
        ws_sum.cell(row=r, column=2).fill = _fill(_LIGHT_GREY if r % 2 == 0 else _WHITE)
        ws_sum.cell(row=r, column=1).border = _thin_border()
        ws_sum.cell(row=r, column=2).border = _thin_border()

    _auto_column_widths(ws_sum)

    # ── Sheet 2: Anomalies ──────────────────────────────────────────────────
    ws_anom = wb.create_sheet("Anomalies")
    anom_headers = [
        "ID", "Detected At", "Employee Code", "Employee Name", "Department",
        "Category", "Severity", "Status", "Baseline Deviation", "Observed Value",
        "Baseline Value", "Description",
    ]
    _apply_header_row(ws_anom, 1, anom_headers)

    anomalies_q = (
        db.execute(
            select(Anomaly, Employee)
            .join(Employee, Employee.id == Anomaly.employee_id)
            .order_by(Anomaly.detected_at.desc())
        )
        .all()
    )
    for r, (anom, emp) in enumerate(anomalies_q, start=2):
        values = [
            str(anom.id),
            anom.detected_at.strftime("%Y-%m-%d %H:%M"),
            emp.employee_code,
            emp.full_name,
            emp.department,
            str(anom.category),
            str(anom.severity),
            str(anom.status),
            round(anom.baseline_deviation, 3),
            round(anom.observed_value, 3),
            round(anom.baseline_value, 3),
            anom.description[:200],
        ]
        # col 7 = severity (1-based), col 8 = status
        color_map = {7: _SEVERITY_HEX.get(anom.severity, _NAVY)}
        _apply_data_row(ws_anom, r, values, alternate=(r % 2 == 0), color_map=color_map)

    _freeze_header(ws_anom)
    _auto_column_widths(ws_anom)

    # ── Sheet 3: Alerts ─────────────────────────────────────────────────────
    ws_alert = wb.create_sheet("Alerts")
    alert_headers = [
        "ID", "Raised At", "Employee Code", "Employee Name", "Department",
        "Title", "Severity", "Status",
    ]
    _apply_header_row(ws_alert, 1, alert_headers)

    alerts_q = (
        db.execute(
            select(Alert, Employee)
            .join(Employee, Employee.id == Alert.employee_id)
            .order_by(Alert.raised_at.desc())
        )
        .all()
    )
    for r, (alert, emp) in enumerate(alerts_q, start=2):
        values = [
            str(alert.id),
            alert.raised_at.strftime("%Y-%m-%d %H:%M"),
            emp.employee_code,
            emp.full_name,
            emp.department,
            alert.title,
            str(alert.severity),
            str(alert.status),
        ]
        color_map = {
            7: _SEVERITY_HEX.get(alert.severity, _NAVY),
            8: _ALERT_STATUS_HEX.get(alert.status, _NAVY),
        }
        _apply_data_row(ws_alert, r, values, alternate=(r % 2 == 0), color_map=color_map)

    _freeze_header(ws_alert)
    _auto_column_widths(ws_alert)

    # ── Sheet 4: Investigations ─────────────────────────────────────────────
    ws_inv = wb.create_sheet("Investigations")
    inv_headers = [
        "Reference", "Title", "Employee Code", "Employee Name", "Department",
        "Severity", "Status", "Created At", "Resolved At", "Assigned To",
    ]
    _apply_header_row(ws_inv, 1, inv_headers)

    invs_q = (
        db.execute(
            select(Investigation, Employee)
            .join(Employee, Employee.id == Investigation.employee_id)
            .order_by(Investigation.created_at.desc())
        )
        .all()
    )
    for r, (inv, emp) in enumerate(invs_q, start=2):
        values = [
            inv.reference,
            inv.title,
            emp.employee_code,
            emp.full_name,
            emp.department,
            str(inv.severity),
            str(inv.status),
            inv.created_at.strftime("%Y-%m-%d %H:%M") if inv.created_at else "",
            inv.resolved_at.strftime("%Y-%m-%d %H:%M") if inv.resolved_at else "Pending",
            inv.assigned_to.full_name if inv.assigned_to else "Unassigned",
        ]
        color_map = {
            6: _SEVERITY_HEX.get(inv.severity, _NAVY),
            7: _INV_STATUS_HEX.get(inv.status, _NAVY),
        }
        _apply_data_row(ws_inv, r, values, alternate=(r % 2 == 0), color_map=color_map)

    _freeze_header(ws_inv)
    _auto_column_widths(ws_inv)

    # ── Sheet 5: Risk Scores ────────────────────────────────────────────────
    ws_risk = wb.create_sheet("Risk Scores")
    risk_headers = [
        "Employee Code", "Employee Name", "Department", "Designation",
        "Risk Score", "Risk Band", "Logon Count", "After-Hours Logons",
        "USB Connects", "File Copies", "Email Count", "Computed At",
    ]
    _apply_header_row(ws_risk, 1, risk_headers)

    lrsq = _latest_rs_subq()
    all_risk = (
        db.execute(
            select(Employee, RiskScore)
            .join(RiskScore, RiskScore.employee_id == Employee.id)
            .join(
                lrsq,
                (lrsq.c.employee_id == RiskScore.employee_id)
                & (lrsq.c.max_computed_at == RiskScore.computed_at),
            )
            .order_by(RiskScore.risk_score.desc())
        )
        .all()
    )
    for r, (emp, rs) in enumerate(all_risk, start=2):
        values = [
            emp.employee_code,
            emp.full_name,
            emp.department,
            emp.designation,
            round(rs.risk_score, 4),
            str(rs.risk_band),
            int(rs.logon_count),
            int(rs.after_hours_logon_count),
            int(rs.usb_connect_count),
            int(rs.file_copy_count),
            int(rs.email_count),
            rs.computed_at.strftime("%Y-%m-%d %H:%M"),
        ]
        color_map = {6: _RISK_HEX.get(rs.risk_band, _NAVY)}
        _apply_data_row(ws_risk, r, values, alternate=(r % 2 == 0), color_map=color_map)

    _freeze_header(ws_risk)
    _auto_column_widths(ws_risk)

    # ── Metadata ────────────────────────────────────────────────────────────
    _add_metadata_sheet(wb, "Insider Threat Report")

    wb.save(str(output_path))
    return output_path


# ---------------------------------------------------------------------------
# 2. Behavioral Analytics Excel Report
# ---------------------------------------------------------------------------


def generate_behavioral_excel(db: Session, output_path: str | Path) -> Path:
    """Multi-sheet behavioral analytics workbook.

    Sheets
    ------
    - Behavior Profiles : Per-employee baselines
    - Activity Logs     : All activity log records
    - Anomaly Detail    : All anomalies with deviation metrics
    - Metadata
    """
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    wb.remove(wb.active)

    # ── Sheet 1: Behavior Profiles ──────────────────────────────────────────
    ws_bp = wb.create_sheet("Behavior Profiles")
    bp_headers = [
        "Employee Code", "Full Name", "Department", "Designation",
        "Login Start (hr)", "Login End (hr)", "Daily Logins", "Daily Downloads",
        "Daily Transfers", "Daily Emails", "Daily Data (MB)", "Typical Devices",
        "Typical Applications",
    ]
    _apply_header_row(ws_bp, 1, bp_headers)

    profiles_q = (
        db.execute(
            select(BehaviorProfile, Employee)
            .join(Employee, Employee.id == BehaviorProfile.employee_id)
            .order_by(Employee.department, Employee.last_name)
        )
        .all()
    )
    for r, (bp, emp) in enumerate(profiles_q, start=2):
        values = [
            emp.employee_code,
            emp.full_name,
            emp.department,
            emp.designation,
            bp.typical_login_hour_start,
            bp.typical_login_hour_end,
            round(bp.typical_daily_logins, 2),
            round(bp.typical_daily_downloads, 2),
            round(bp.typical_daily_transfers, 2),
            round(bp.typical_daily_emails, 2),
            round(bp.typical_daily_data_volume_mb, 2),
            bp.typical_device_count,
            bp.typical_applications[:100] if bp.typical_applications else "",
        ]
        _apply_data_row(ws_bp, r, values, alternate=(r % 2 == 0))

    _freeze_header(ws_bp)
    _auto_column_widths(ws_bp)

    # ── Sheet 2: Activity Logs ──────────────────────────────────────────────
    ws_act = wb.create_sheet("Activity Logs")
    act_headers = [
        "ID", "Timestamp", "Employee Code", "Employee Name", "Department",
        "Activity Type", "Source", "Device", "IP Address", "Application",
        "Data Volume (MB)", "Details",
    ]
    _apply_header_row(ws_act, 1, act_headers)

    activities_q = (
        db.execute(
            select(ActivityLog, Employee)
            .join(Employee, Employee.id == ActivityLog.employee_id)
            .order_by(ActivityLog.timestamp.desc())
        )
        .all()
    )
    for r, (act, emp) in enumerate(activities_q, start=2):
        values = [
            str(act.id),
            act.timestamp.strftime("%Y-%m-%d %H:%M"),
            emp.employee_code,
            emp.full_name,
            emp.department,
            str(act.activity_type),
            act.source,
            act.device,
            act.ip_address,
            act.application or "",
            round(act.data_volume_mb, 3),
            (act.details or "")[:150],
        ]
        _apply_data_row(ws_act, r, values, alternate=(r % 2 == 0))

    _freeze_header(ws_act)
    _auto_column_widths(ws_act)

    # ── Sheet 3: Anomaly Detail ─────────────────────────────────────────────
    ws_anom = wb.create_sheet("Anomaly Detail")
    anom_headers = [
        "ID", "Detected At", "Employee Code", "Employee Name", "Department",
        "Category", "Severity", "Status",
        "Baseline Deviation", "Observed Value", "Baseline Value",
        "Description",
    ]
    _apply_header_row(ws_anom, 1, anom_headers)

    anomalies_q = (
        db.execute(
            select(Anomaly, Employee)
            .join(Employee, Employee.id == Anomaly.employee_id)
            .order_by(Anomaly.baseline_deviation.desc())
        )
        .all()
    )
    for r, (anom, emp) in enumerate(anomalies_q, start=2):
        values = [
            str(anom.id),
            anom.detected_at.strftime("%Y-%m-%d %H:%M"),
            emp.employee_code,
            emp.full_name,
            emp.department,
            str(anom.category),
            str(anom.severity),
            str(anom.status),
            round(anom.baseline_deviation, 3),
            round(anom.observed_value, 3),
            round(anom.baseline_value, 3),
            anom.description[:200],
        ]
        color_map = {7: _SEVERITY_HEX.get(anom.severity, _NAVY)}
        _apply_data_row(ws_anom, r, values, alternate=(r % 2 == 0), color_map=color_map)

    _freeze_header(ws_anom)
    _auto_column_widths(ws_anom)

    _add_metadata_sheet(wb, "Behavioral Analytics Report")

    wb.save(str(output_path))
    return output_path


# ---------------------------------------------------------------------------
# 3. Risk Assessment Excel Report
# ---------------------------------------------------------------------------


def generate_risk_assessment_excel(db: Session, output_path: str | Path) -> Path:
    """Multi-sheet risk assessment workbook.

    Sheets
    ------
    - Risk Summary      : Band distribution + KPIs
    - Employee Risks    : Full risk score table (latest per employee)
    - Risk History      : All risk score records across time
    - Metadata
    """
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    wb.remove(wb.active)

    lrsq = _latest_rs_subq()

    # ── Sheet 1: Risk Summary ───────────────────────────────────────────────
    ws_sum = wb.create_sheet("Risk Summary")
    _apply_header_row(ws_sum, 1, ["Metric", "Value"])

    assessed_employees = (
        db.scalar(
            select(func.count())
            .select_from(RiskScore)
            .join(
                lrsq,
                (lrsq.c.employee_id == RiskScore.employee_id)
                & (lrsq.c.max_computed_at == RiskScore.computed_at),
            )
        )
        or 0
    )
    avg_score = (
        db.scalar(
            select(func.avg(RiskScore.risk_score))
            .select_from(RiskScore)
            .join(
                lrsq,
                (lrsq.c.employee_id == RiskScore.employee_id)
                & (lrsq.c.max_computed_at == RiskScore.computed_at),
            )
        )
    )

    kpi_rows: list[tuple[str, object]] = [
        ("Assessed Employees", assessed_employees),
        ("Average Risk Score", round(avg_score, 4) if avg_score is not None else "N/A"),
    ]
    for band in RiskBand:
        count = (
            db.scalar(
                select(func.count())
                .select_from(RiskScore)
                .join(
                    lrsq,
                    (lrsq.c.employee_id == RiskScore.employee_id)
                    & (lrsq.c.max_computed_at == RiskScore.computed_at),
                )
                .where(RiskScore.risk_band == band)
            )
            or 0
        )
        kpi_rows.append((f"{str(band)} Risk Employees", count))

    kpi_rows.append(("Report Generated", _now_str()))

    for r, (metric, val) in enumerate(kpi_rows, start=2):
        ws_sum.cell(row=r, column=1, value=metric).font = _body_font(bold=True)
        ws_sum.cell(row=r, column=2, value=val).font = _body_font()
        ws_sum.cell(row=r, column=1).fill = _fill(_LIGHT_GREY if r % 2 == 0 else _WHITE)
        ws_sum.cell(row=r, column=2).fill = _fill(_LIGHT_GREY if r % 2 == 0 else _WHITE)
        ws_sum.cell(row=r, column=1).border = _thin_border()
        ws_sum.cell(row=r, column=2).border = _thin_border()

    _auto_column_widths(ws_sum)

    # ── Sheet 2: Employee Risks (Latest) ────────────────────────────────────
    ws_emp = wb.create_sheet("Employee Risks")
    emp_headers = [
        "Employee Code", "Full Name", "Department", "Designation",
        "Risk Score", "Risk Band", "Predict Label",
        "Decision Function Score", "Logon Count", "After-Hours Logons",
        "USB Connects", "File Copies", "Email Count",
        "Lookback Days", "Computed At",
    ]
    _apply_header_row(ws_emp, 1, emp_headers)

    all_latest = (
        db.execute(
            select(Employee, RiskScore)
            .join(RiskScore, RiskScore.employee_id == Employee.id)
            .join(
                lrsq,
                (lrsq.c.employee_id == RiskScore.employee_id)
                & (lrsq.c.max_computed_at == RiskScore.computed_at),
            )
            .order_by(RiskScore.risk_score.desc())
        )
        .all()
    )
    for r, (emp, rs) in enumerate(all_latest, start=2):
        values = [
            emp.employee_code,
            emp.full_name,
            emp.department,
            emp.designation,
            round(rs.risk_score, 4),
            str(rs.risk_band),
            rs.predict_label,
            round(rs.decision_function_score, 4),
            int(rs.logon_count),
            int(rs.after_hours_logon_count),
            int(rs.usb_connect_count),
            int(rs.file_copy_count),
            int(rs.email_count),
            rs.lookback_days,
            rs.computed_at.strftime("%Y-%m-%d %H:%M"),
        ]
        color_map = {6: _RISK_HEX.get(rs.risk_band, _NAVY)}
        _apply_data_row(ws_emp, r, values, alternate=(r % 2 == 0), color_map=color_map)

    _freeze_header(ws_emp)
    _auto_column_widths(ws_emp)

    # ── Sheet 3: Risk History (all records) ─────────────────────────────────
    ws_hist = wb.create_sheet("Risk History")
    hist_headers = [
        "Employee Code", "Full Name", "Department",
        "Risk Score", "Risk Band", "Predict Label",
        "Logon Count", "After-Hours Logons", "USB Connects",
        "File Copies", "Email Count", "Lookback Days", "Computed At",
    ]
    _apply_header_row(ws_hist, 1, hist_headers)

    all_history = (
        db.execute(
            select(Employee, RiskScore)
            .join(RiskScore, RiskScore.employee_id == Employee.id)
            .order_by(RiskScore.computed_at.desc())
        )
        .all()
    )
    for r, (emp, rs) in enumerate(all_history, start=2):
        values = [
            emp.employee_code,
            emp.full_name,
            emp.department,
            round(rs.risk_score, 4),
            str(rs.risk_band),
            rs.predict_label,
            int(rs.logon_count),
            int(rs.after_hours_logon_count),
            int(rs.usb_connect_count),
            int(rs.file_copy_count),
            int(rs.email_count),
            rs.lookback_days,
            rs.computed_at.strftime("%Y-%m-%d %H:%M"),
        ]
        color_map = {5: _RISK_HEX.get(rs.risk_band, _NAVY)}
        _apply_data_row(ws_hist, r, values, alternate=(r % 2 == 0), color_map=color_map)

    _freeze_header(ws_hist)
    _auto_column_widths(ws_hist)

    _add_metadata_sheet(wb, "Risk Assessment Report")

    wb.save(str(output_path))
    return output_path
