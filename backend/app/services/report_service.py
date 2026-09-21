import io
import csv
from datetime import datetime
from typing import List, Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

def generate_csv_report(data: List[Dict[str, Any]], fieldnames: List[str]) -> str:
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=fieldnames, extrasaction="ignore")
    writer.writeheader()
    for row in data:
        writer.writerow(row)
    return output.getvalue()

def generate_excel_report(report_title: str, sheets_data: Dict[str, List[Dict[str, Any]]]) -> bytes:
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)
    
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    
    for sheet_name, rows in sheets_data.items():
        ws = wb.create_sheet(title=sheet_name[:30])
        if not rows:
            ws.append(["No Data Available"])
            continue
            
        headers = list(rows[0].keys())
        ws.append(headers)
        
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
        for row in rows:
            ws.append([str(row.get(h, "")) for h in headers])
            
        for row_cells in ws.iter_rows(min_row=1, max_row=len(rows)+1, min_col=1, max_col=len(headers)):
            for cell in row_cells:
                cell.border = thin_border
                
        # Auto-fit column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
            
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()

def generate_pdf_report(report_title: str, subtitle: str, summary_stats: Dict[str, Any], table_data: List[List[str]]) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'ReportTitle',
        parent=styles['Heading1'],
        fontSize=20,
        leading=24,
        textColor=colors.HexColor("#0f172a"),
        spaceAfter=6
    )
    subtitle_style = ParagraphStyle(
        'ReportSubtitle',
        parent=styles['Normal'],
        fontSize=10,
        textColor=colors.HexColor("#64748b"),
        spaceAfter=14
    )
    section_style = ParagraphStyle(
        'SectionHeader',
        parent=styles['Heading2'],
        fontSize=13,
        leading=16,
        textColor=colors.HexColor("#1e293b"),
        spaceBefore=10,
        spaceAfter=8
    )
    body_style = styles['Normal']
    
    story = []
    
    # Title & Subtitle
    story.append(Paragraph(report_title, title_style))
    story.append(Paragraph(f"{subtitle} | Generated on {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}", subtitle_style))
    story.append(Spacer(1, 10))
    
    # Summary Table
    if summary_stats:
        story.append(Paragraph("Executive Summary & Risk Metrics", section_style))
        sum_data = [["Metric", "Value"]] + [[k.replace('_', ' ').title(), str(v)] for k, v in summary_stats.items()]
        t_summary = Table(sum_data, colWidths=[240, 280])
        t_summary.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1e293b")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor("#f8fafc"), colors.white]),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
        ]))
        story.append(t_summary)
        story.append(Spacer(1, 14))
        
    # Main Details Table
    if table_data and len(table_data) > 1:
        story.append(Paragraph("Detailed Incident & Behavioral Records", section_style))
        t_main = Table(table_data, colWidths=[int(520 / len(table_data[0]))] * len(table_data[0]))
        t_main.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#334155")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 5),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f5f9")]),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ]))
        story.append(t_main)
        
    doc.build(story)
    return buffer.getvalue()
